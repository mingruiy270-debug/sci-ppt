"""Build editable academic slides from an explicit JSON specification."""
import argparse
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


def color(value):
    return RGBColor.from_string(value.lstrip('#'))


def build(spec_path, output):
    spec_path = Path(spec_path).resolve()
    spec = json.loads(spec_path.read_text(encoding='utf-8-sig'))
    prs = Presentation()
    sw, sh = spec.get('size', [13.333333, 7.5])
    prs.slide_width, prs.slide_height = Inches(sw), Inches(sh)
    theme = {'font': 'Arial', 'ink': '20292C', 'accent': '167D8D',
             'muted': '526568', 'paper': 'FFFFFF'}
    theme.update(spec.get('theme', {}))
    prs.core_properties.title = spec.get('title', '')
    prs.core_properties.author = spec.get('author', '')

    def text(slide, value, box, size=22, bold=False, fill=None, name='text'):
        x, y, w, h = box
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.025)
        tf.margin_top = tf.margin_bottom = Inches(0.015)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(value.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.space_after = Pt(5)
            for run in p.runs:
                run.font.name = theme['font']
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = color(fill or theme['ink'])
                ea = OxmlElement('a:ea')
                ea.set('typeface', theme['font'])
                run._r.get_or_add_rPr().append(ea)
        return shape

    def rect(slide, box, fill):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *[Inches(v) for v in box])
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill)
        shape.line.fill.background()
        shape._element.spPr.append(OxmlElement('a:effectLst'))
        return shape

    def picture(slide, item):
        path = (spec_path.parent / item['path']).resolve()
        if not path.is_file():
            raise FileNotFoundError(path)
        with Image.open(path) as im:
            iw, ih = im.size
        x, y, w, h = item['box']
        scale = min(w / iw, h / ih)
        aw, ah = iw * scale, ih * scale
        shape = slide.shapes.add_picture(str(path), Inches(x+(w-aw)/2),
                                        Inches(y+(h-ah)/2), Inches(aw), Inches(ah))
        shape.name = item.get('id', path.stem)

    for idx, page in enumerate(spec['slides'], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color(theme['paper'])
        rect(slide, [0, 0, 0.11, sh], theme['accent'])
        text(slide, page.get('section', ''), [0.48, 0.22, 11.8, 0.28], 12,
             fill=theme['accent'], name='section')
        text(slide, page['title'], page.get('title_box', [0.48, 0.64, 12.2, 0.8]),
             page.get('title_size', 30), True, name='title')
        for item in page.get('shapes', []):
            if item.get('kind') == 'arrow':
                shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                               *[Inches(v) for v in item['box']])
                shape.fill.solid()
                shape.fill.fore_color.rgb = color(item.get('fill', theme['accent']))
                shape.line.fill.background()
                shape._element.spPr.append(OxmlElement('a:effectLst'))
            else:
                rect(slide, item['box'], item.get('fill', 'F0F5F5'))
        for item in page.get('images', []):
            picture(slide, item)
        for item in page.get('texts', []):
            text(slide, item['text'], item['box'], item.get('size', 22),
                 item.get('bold', False), item.get('color'), item.get('id', 'body'))
        takeaway = page.get('takeaway')
        if takeaway:
            rect(slide, [0.48, 6.58, 12.35, 0.48], 'EDF5F5')
            text(slide, takeaway, [0.62, 6.63, 12.03, 0.36], 17, True,
                 fill=theme['accent'], name='takeaway')
        text(slide, page.get('footer', page.get('source', '')), [0.48, 7.14, 11.8, 0.26], 10,
             fill=theme['muted'], name='source')
        text(slide, f'{idx:02d}', [12.3, 7.14, 0.5, 0.25], 11,
             fill=theme['muted'], name='page')
        notes = page.get('notes', '')
        if page.get('source'):
            notes += '\n\nSource: ' + page['source']
        slide.notes_slide.notes_text_frame.text = notes
        for shape in slide.shapes:
            if min(shape.left, shape.top) < 0 or shape.left+shape.width > prs.slide_width+10 or shape.top+shape.height > prs.slide_height+10:
                raise ValueError(f'Slide {idx}: out-of-bounds shape {shape.name}')
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    reopened = Presentation(output)
    if len(reopened.slides) != len(spec['slides']):
        raise ValueError('Slide count changed on reopen')
    print(f'Saved {output}; {len(reopened.slides)} slides with editable text and notes.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('spec', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(args.spec, args.output)
